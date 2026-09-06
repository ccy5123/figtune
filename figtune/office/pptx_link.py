"""PowerPoint 연동.

IguanaTeX의 구조를 따른다: 그림과 함께 **그림을 만든 소스를 도형에 심어두고**,
나중에 그 도형을 골라 다시 편집·재생성한다. 슬라이드는 소스를 품고 다니므로
figtune이 없는 컴퓨터에서도 발표는 되고, 편집만 figtune을 요구한다.

저장 위치는 도형의 alt text다. VBA의 Shape.Tags보다 접근성이 낫고
(사람이 읽을 설명을 앞에 둘 수 있다) python-pptx와 VBA 양쪽에서 읽고 쓸 수
있어 PowerPoint 없이도 같은 파일을 다룰 수 있다.

payload 구성:
  script — 원본 플로팅 스크립트 경로 (deck 기준 상대경로 우선)
  spec   — figtune spec 전체를 zlib 압축 후 base64 (보통 1KB 미만)

스크립트와 spec을 둘 다 넣는 이유: 스크립트가 있으면 데이터를 다시 돌려
그림을 갱신할 수 있고, 없으면 최소한 어떤 스타일이었는지는 남는다.
"""

from __future__ import annotations

import base64
import json
import os
import re
import zlib
from dataclasses import dataclass
from pathlib import Path

import yaml

from ..core.spec import Spec
from ..i18n import t as _t

MARK_RE = re.compile(r"<figtune:(\d+)>([A-Za-z0-9+/=]+)</figtune>", re.S)
PAYLOAD_VERSION = 1


@dataclass
class Payload:
    script: str
    spec: Spec
    dpi: int = 300
    description: str = ""
    # 벡터로 넣을지. PNG 대체본이 항상 함께 들어가므로 공유는 안전하다.
    vector: bool = False

    def resolve_script(self, deck_path: Path) -> Path:
        p = Path(self.script)
        return p if p.is_absolute() else (Path(deck_path).parent / p).resolve()

    @classmethod
    def for_deck(cls, script, deck, spec, dpi: int = 300, description: str = "",
                 vector: bool = False):
        """덱 기준 상대경로로 payload를 만든다.

        기준점이 둘이라 헷갈리기 쉽다. **스크립트 경로는 덱 기준**이고,
        **스크립트가 읽는 데이터 경로는 스크립트 기준**이다. 후자는 실행 시
        스크립트 디렉토리로 chdir하므로 자동으로 맞는다.

        상대경로로 만들 수 없으면(다른 드라이브 등) 절대경로로 남긴다.
        """
        script, deck = Path(script).resolve(), Path(deck).resolve()
        try:
            rel = os.path.relpath(script, deck.parent)
        except ValueError:
            rel = str(script)
        return cls(script=rel, spec=spec, dpi=dpi, description=description,
                   vector=vector)


def pack(payload: Payload) -> str:
    """alt text 문자열을 만든다. 앞부분은 사람이 읽는 설명."""
    body = json.dumps(
        {"script": payload.script, "dpi": payload.dpi,
         "vector": payload.vector,
         "spec": yaml.safe_dump(payload.spec.to_dict(), allow_unicode=True)},
        ensure_ascii=False).encode("utf-8")
    blob = base64.b64encode(zlib.compress(body, 9)).decode("ascii")
    human = payload.description or f"figtune figure ({Path(payload.script).name})"
    return f"{human}\n\n<figtune:{PAYLOAD_VERSION}>{blob}</figtune>"


def unpack(alt_text: str | None) -> Payload | None:
    """alt text에서 payload를 꺼낸다. figtune 도형이 아니면 None."""
    if not alt_text:
        return None
    m = MARK_RE.search(alt_text)
    if not m:
        return None
    if int(m.group(1)) > PAYLOAD_VERSION:
        raise ValueError(_t(
            "payload 버전 {v}은 이 figtune보다 새롭습니다. 업데이트하세요.",
            v=m.group(1)))
    try:
        d = json.loads(zlib.decompress(base64.b64decode(m.group(2))).decode("utf-8"))
    except Exception as exc:
        raise ValueError(_t("figtune payload를 읽지 못했습니다: {err}",
                            err=exc)) from exc
    return Payload(script=d["script"], spec=Spec.from_dict(yaml.safe_load(d["spec"])),
                   dpi=int(d.get("dpi", 300)),
                   vector=bool(d.get("vector", False)),
                   description=alt_text.split("\n\n<figtune:")[0])


# --- 도형 조작 ------------------------------------------------------------

def _alt(shape) -> str | None:
    try:
        return shape._element._nvXxPr.cNvPr.get("descr")
    except AttributeError:
        return None


def _set_alt(shape, text: str) -> None:
    shape._element._nvXxPr.cNvPr.set("descr", text)


def iter_figures(prs):
    """프레젠테이션 안의 figtune 도형을 (slide_index, shape, payload)로 훑는다."""
    for i, slide in enumerate(prs.slides):
        for shape in list(slide.shapes):
            try:
                payload = unpack(_alt(shape))
            except ValueError:
                payload = None
            if payload is not None:
                yield i, shape, payload


def insert(slide, image_path, payload: Payload, left, top,
           width=None, height=None):
    """슬라이드에 그림을 넣고 payload를 심는다."""
    pic = slide.shapes.add_picture(str(image_path), left, top,
                                   width=width, height=height)
    _set_alt(pic, pack(payload))
    return pic


def replace(slide, shape, image_path, payload: Payload | None = None,
            svg_path=None):
    """기존 도형을 새 그림으로 갈아끼운다.

    위치·크기·회전·z순서를 보존한다. 발표자가 슬라이드에서 손으로 맞춰둔
    배치를 재생성 때문에 잃으면 도구를 쓸 이유가 없다.
    """
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    rotation = getattr(shape, "rotation", 0)
    alt = pack(payload) if payload is not None else _alt(shape)

    if svg_path is not None:
        from .vector import add_svg_picture
        pic = add_svg_picture(slide, svg_path, image_path, left, top,
                              width=width, height=height)
    else:
        pic = slide.shapes.add_picture(str(image_path), left, top,
                                       width=width, height=height)
    try:
        pic.rotation = rotation
    except (AttributeError, ValueError):
        pass
    if alt:
        _set_alt(pic, alt)

    old_el, new_el = shape._element, pic._element
    old_el.addprevious(new_el)          # z순서 유지
    old_el.getparent().remove(old_el)
    return pic


# --- 렌더링 / 갱신 --------------------------------------------------------

def render(script: Path | str, spec: Spec, out_path: Path | str,
           dpi: int = 300, python: str | None = None):
    """스크립트를 실행하고 spec을 입혀 이미지로 뽑는다.

    (경로, 데이터 변화)를 돌려준다. 데이터 변화는 이 그림이 만들어질 당시와
    지금의 입력 파일이 같은지를 말해준다.
    """
    import matplotlib
    matplotlib.use("Agg")
    from ..core.session import Session

    s = Session(python=python)
    s.open(script, spec=spec)
    changes = s.data_changes()
    return s.export(out_path, dpi=dpi), changes


@dataclass
class RefreshResult:
    updated: list[str]
    skipped: list[tuple[str, str]]      # (설명, 사유)
    data_changes: dict = None           # 설명 -> provenance.compare 결과

    def __post_init__(self):
        if self.data_changes is None:
            self.data_changes = {}

    def summary(self) -> str:
        out = [_t("갱신 {n}건", n=len(self.updated))]
        if self.skipped:
            out.append(_t("건너뜀 {n}건", n=len(self.skipped)))
        n = sum(1 for d in self.data_changes.values()
                if d["changed"] or d["added"] or d["removed"])
        if n:
            out.append(_t("데이터 변경 {n}건", n=n))
        return ", ".join(out)


def refresh_deck(pptx_path: Path | str, out_path: Path | str | None = None,
                 python: str | None = None, check_only: bool = False) -> RefreshResult:
    """덱 안의 모든 figtune 그림을 원본 스크립트로 다시 만든다.

    데이터나 모델이 바뀌었을 때 슬라이드 전체를 한 번에 최신화한다.
    figtune이 PowerPoint에 붙어야 하는 진짜 이유가 이것이다.
    """
    import tempfile
    from pptx import Presentation

    src = Path(pptx_path)
    prs = Presentation(str(src))
    updated, skipped, changes = [], [], {}

    with tempfile.TemporaryDirectory() as td:
        for idx, shape, payload in list(iter_figures(prs)):
            label = f"slide {idx + 1} / {Path(payload.script).name}"
            script = payload.resolve_script(src)
            if not script.exists():
                skipped.append((label, _t("스크립트를 찾을 수 없음: {path}",
                                          path=script)))
                continue
            stem = Path(td) / f"fig_{idx}_{id(shape)}"
            svg = None
            try:
                if payload.vector:
                    from .vector import render_pair
                    svg, png, diff = render_pair(script, payload.spec, stem,
                                                 dpi=payload.dpi, python=python)
                else:
                    png = stem.with_suffix(".png")
                    _, diff = render(script, payload.spec, png,
                                     dpi=payload.dpi, python=python)
            except Exception as exc:
                skipped.append((label, _t("렌더 실패: {err}", err=exc)))
                continue
            changes[label] = diff
            if not check_only:
                replace(prs.slides[idx], shape, png, payload, svg_path=svg)
            updated.append(label)

        if not check_only:
            prs.save(str(out_path or src))

    return RefreshResult(updated, skipped, changes)
