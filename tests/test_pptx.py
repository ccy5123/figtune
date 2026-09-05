"""PowerPoint 연동 테스트.

IguanaTeX 구조의 핵심 두 가지를 확인한다.
  1. 슬라이드가 소스를 품고 다닌다 (저장 → 다시 열기 → 편집 가능)
  2. 재생성이 발표자의 배치를 망가뜨리지 않는다
"""

import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")

import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

from figtune.core.session import Session
from figtune.office import pptx_link as PL

EXAMPLE = Path(__file__).parent.parent / "examples" / "plot_fig3.py"


@pytest.fixture
def script(tmp_path):
    dst = tmp_path / "plot_fig3.py"
    shutil.copy(EXAMPLE, dst)
    return dst


@pytest.fixture
def styled_spec(script):
    s = Session()
    s.open(script)
    s.set_prop("ax0.line0", "color", "#c0392b")
    s.set_prop("ax0.spine:top", "visible", False)
    s.set_prop("ax0.xtick.major", "locator", {"kind": "multiple", "base": 12.0})
    s.add_panel_labels()
    return s.spec


def test_payload_survives_pack_unpack(styled_spec, script):
    p = PL.Payload(script=script.name, spec=styled_spec, dpi=300)
    got = PL.unpack(PL.pack(p))
    assert got is not None
    assert got.script == script.name
    assert got.dpi == 300
    assert got.spec.overrides == styled_spec.overrides
    assert [t.text for t in got.spec.texts] == [t.text for t in styled_spec.texts]


def test_payload_stays_small_enough_for_alt_text(styled_spec, script):
    alt = PL.pack(PL.Payload(script=script.name, spec=styled_spec))
    assert len(alt) < 4000, f"alt text가 너무 큽니다: {len(alt)}"


def test_non_figtune_shape_is_ignored():
    assert PL.unpack(None) is None
    assert PL.unpack("그냥 사진입니다") is None


def test_deck_carries_source_and_survives_reopen(tmp_path, script, styled_spec):
    png = tmp_path / "fig.png"
    PL.render(script, styled_spec, png, dpi=150)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    PL.insert(slide, png, PL.Payload(script=script.name, spec=styled_spec, dpi=150),
              left=Inches(0.5), top=Inches(1.0), width=Inches(8))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    # 다시 열었을 때 소스가 살아 있어야 편집이 가능하다
    found = list(PL.iter_figures(Presentation(str(deck))))
    assert len(found) == 1
    _, _, payload = found[0]
    assert payload.spec.get("ax0.line0", "color") == "#c0392b"
    assert payload.resolve_script(deck) == script


def test_refresh_preserves_placement_and_zorder(tmp_path, script, styled_spec):
    png = tmp_path / "fig.png"
    PL.render(script, styled_spec, png, dpi=100)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # figtune 그림 앞뒤로 다른 도형을 둬서 z순서를 확인한다
    slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(2), Inches(0.5))
    PL.insert(slide, png, PL.Payload(script=script.name, spec=styled_spec, dpi=100),
              left=Inches(1.25), top=Inches(2.0), width=Inches(6))
    slide.shapes.add_textbox(Inches(0.2), Inches(6.0), Inches(2), Inches(0.5))

    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    before = Presentation(str(deck))
    _, shp, _ = next(iter(PL.iter_figures(before)))
    pos = (shp.left, shp.top, shp.width, shp.height)
    index_before = list(before.slides[0].shapes).index(shp)

    result = PL.refresh_deck(deck)
    assert result.updated and not result.skipped, result.skipped

    after = Presentation(str(deck))
    _, shp2, _ = next(iter(PL.iter_figures(after)))
    assert (shp2.left, shp2.top, shp2.width, shp2.height) == pos
    assert list(after.slides[0].shapes).index(shp2) == index_before
    assert len(list(after.slides[0].shapes)) == 3


def test_refresh_reports_missing_script_instead_of_crashing(tmp_path, script,
                                                            styled_spec):
    png = tmp_path / "fig.png"
    PL.render(script, styled_spec, png, dpi=100)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    PL.insert(slide, png,
              PL.Payload(script="없는파일.py", spec=styled_spec, dpi=100),
              left=Inches(1), top=Inches(1), width=Inches(6))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    result = PL.refresh_deck(deck)
    assert not result.updated
    assert len(result.skipped) == 1
    assert "찾을 수 없음" in result.skipped[0][1]


def test_refresh_picks_up_changed_data(tmp_path, script, styled_spec):
    """데이터가 바뀌면 덱이 실제로 갱신되어야 한다."""
    png = tmp_path / "fig.png"
    PL.render(script, styled_spec, png, dpi=100)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    PL.insert(slide, png, PL.Payload(script=script.name, spec=styled_spec, dpi=100),
              left=Inches(1), top=Inches(1), width=Inches(6))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))

    def image_bytes(path):
        _, shp, _ = next(iter(PL.iter_figures(Presentation(str(path)))))
        return shp.image.blob

    original = image_bytes(deck)

    # 원본 스크립트의 데이터를 바꾼다 (모델 재실행 상황)
    script.write_text(script.read_text().replace("0.09 * t", "0.20 * t"),
                      encoding="utf-8")
    assert PL.refresh_deck(deck).updated
    assert image_bytes(deck) != original


def test_vba_bridge_roundtrip(tmp_path, script, styled_spec):
    """VBA는 zlib이 없어 인코딩을 파이썬에 위임한다. 그 다리가 왕복해야 한다."""
    from figtune.office import vba_bridge

    spec_file = tmp_path / "in.yaml"
    styled_spec.dump(spec_file)
    alt_file = tmp_path / "alt.txt"

    assert vba_bridge.main(["pack", "--script", str(script),
                            "--spec", str(spec_file), "--dpi", "200",
                            "--out", str(alt_file)]) == 0

    spec_out, script_out = tmp_path / "out.yaml", tmp_path / "script.txt"
    assert vba_bridge.main(["unpack", "--alt", str(alt_file),
                            "--spec-out", str(spec_out),
                            "--script-out", str(script_out)]) == 0

    from figtune.core.spec import Spec
    assert Spec.load(spec_out).overrides == styled_spec.overrides
    assert script_out.read_text(encoding="utf-16").strip() == str(script)


def test_vba_bridge_rejects_non_figtune_alt(tmp_path):
    from figtune.office import vba_bridge
    alt = tmp_path / "alt.txt"
    alt.write_text("보통 사진입니다", encoding="utf-16")
    assert vba_bridge.main(["unpack", "--alt", str(alt),
                            "--spec-out", str(tmp_path / "x.yaml")]) == 1


def test_check_mode_reports_without_writing(tmp_path, script, styled_spec):
    """--check는 덱을 건드리지 않고 무엇이 바뀔지만 알려야 한다."""
    png = tmp_path / "fig.png"
    PL.render(script, styled_spec, png, dpi=100)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    PL.insert(slide, png, PL.Payload(script=script.name, spec=styled_spec, dpi=100),
              left=Inches(1), top=Inches(1), width=Inches(6))
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    before = deck.read_bytes()

    result = PL.refresh_deck(deck, check_only=True)
    assert result.updated
    assert deck.read_bytes() == before, "check 모드가 파일을 수정했습니다"


def _project(tmp_path):
    """slides/ 와 analysis/ 가 분리된 전형적인 프로젝트 구조."""
    (tmp_path / "slides").mkdir()
    (tmp_path / "analysis").mkdir()
    (tmp_path / "analysis" / "data.csv").write_text(
        "t,y\n0,1\n1,2\n2,4\n", encoding="utf-8")
    script = tmp_path / "analysis" / "plot_a.py"
    script.write_text(
        "import pandas as pd, matplotlib.pyplot as plt\n"
        "df = pd.read_csv('data.csv')\n"          # 스크립트 기준 상대경로
        "fig, ax = plt.subplots()\n"
        "ax.plot(df.t, df.y, label='obs')\n"
        "ax.legend()\n", encoding="utf-8")
    return script, tmp_path / "slides" / "deck.pptx"


def _build_deck(script, deck):
    s = Session()
    s.open(script)
    s.set_prop("ax0.line0", "color", "#c0392b")
    s.save()
    png = deck.parent / "_tmp.png"
    PL.render(script, s.spec, png, dpi=100)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    PL.insert(slide, png, PL.Payload.for_deck(script, deck, s.spec, dpi=100),
              left=Inches(1), top=Inches(1), width=Inches(5))
    prs.save(str(deck))
    png.unlink()
    return s.spec


def test_script_path_is_relative_to_the_deck(tmp_path):
    """스크립트 경로는 덱 기준, 데이터 경로는 스크립트 기준."""
    script, deck = _project(tmp_path)
    _build_deck(script, deck)

    _, _, payload = next(iter(PL.iter_figures(Presentation(str(deck)))))
    assert not Path(payload.script).is_absolute()
    assert payload.resolve_script(deck) == script.resolve()
    # 데이터는 스크립트 기준 상대경로로 기록된다
    assert [d["path"] for d in payload.spec.data_sources] == ["data.csv"]

    result = PL.refresh_deck(deck)
    assert result.updated and not result.skipped


def test_moving_the_whole_project_is_not_a_data_change(tmp_path):
    """폴더째 옮긴 것을 데이터 변경으로 보고하면 경고가 무의미해진다."""
    script, deck = _project(tmp_path)
    _build_deck(script, deck)

    moved = tmp_path.parent / (tmp_path.name + "_moved")
    shutil.copytree(tmp_path, moved)

    result = PL.refresh_deck(moved / "slides" / "deck.pptx", check_only=True)
    assert result.updated and not result.skipped
    diff = list(result.data_changes.values())[0]
    assert not diff["changed"] and not diff["added"] and not diff["removed"]


def test_real_data_edit_still_reported_after_move(tmp_path):
    script, deck = _project(tmp_path)
    _build_deck(script, deck)
    moved = tmp_path.parent / (tmp_path.name + "_moved2")
    shutil.copytree(tmp_path, moved)

    (moved / "analysis" / "data.csv").write_text(
        "t,y\n0,9\n1,8\n2,7\n", encoding="utf-8")
    result = PL.refresh_deck(moved / "slides" / "deck.pptx", check_only=True)
    assert list(result.data_changes.values())[0]["changed"] == ["data.csv"]


def test_deck_is_self_contained_for_viewers_without_figtune(tmp_path):
    """공유가 되려면 슬라이드에 진짜 그림이 박혀 있어야 한다.

    figtune은 spec을 alt text에 '메타데이터로만' 얹는다. 그림 자체는 보통
    PNG이며 pptx 안에 임베드된다. 외부 링크로 바뀌면 받는 사람 화면에서
    그림이 깨지므로 잠가둔다.
    """
    import zipfile

    script, deck = _project(tmp_path)
    _build_deck(script, deck)

    with zipfile.ZipFile(deck) as z:
        names = z.namelist()
        media = [n for n in names if n.startswith("ppt/media/")]
        assert media, "슬라이드에 이미지가 임베드되지 않았습니다"
        assert any(n.endswith(".png") for n in media)
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()

    assert 'TargetMode="External"' not in rels, "이미지가 외부 링크입니다"
    # 그림 도형이지 OLE 객체가 아니다
    _, shape, _ = next(iter(PL.iter_figures(Presentation(str(deck)))))
    assert shape.shape_type is not None
    assert shape.image.blob[:4] == b"\x89PNG"


def test_vector_mode_embeds_svg_with_png_fallback(tmp_path):
    """SVG를 모르는 뷰어는 PNG를, 아는 뷰어는 벡터를 본다."""
    import zipfile

    from figtune.office import vector as V

    script, deck = _project(tmp_path)
    s = Session()
    s.open(script)
    s.set_prop("ax0.line0", "color", "#c0392b")
    s.save()

    svg, png, _ = V.render_pair(script, s.spec, deck.parent / "pair", dpi=150)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = V.add_svg_picture(slide, svg, png, Inches(1), Inches(1),
                            width=Inches(5))
    payload = PL.Payload.for_deck(script, deck, s.spec, dpi=150, vector=True)
    pic._element._nvXxPr.cNvPr.set("descr", PL.pack(payload))
    prs.save(str(deck))

    with zipfile.ZipFile(deck) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        ctypes = z.read("[Content_Types].xml").decode()
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()

    assert any(n.endswith(".svg") for n in media)
    assert any(n.endswith(".png") for n in media), "PNG 대체본이 없습니다"
    assert "image/svg+xml" in ctypes
    assert 'TargetMode="External"' not in rels

    again = Presentation(str(deck))
    shp = again.slides[0].shapes[0]
    assert V.has_svg(shp)
    got = PL.unpack(shp._element._nvXxPr.cNvPr.get("descr"))
    assert got.vector is True


def test_vector_payload_survives_refresh(tmp_path):
    from figtune.office import vector as V

    script, deck = _project(tmp_path)
    s = Session()
    s.open(script)
    s.save()
    svg, png, _ = V.render_pair(script, s.spec, deck.parent / "pair", dpi=120)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = V.add_svg_picture(slide, svg, png, Inches(1), Inches(1), width=Inches(5))
    pic._element._nvXxPr.cNvPr.set(
        "descr", PL.pack(PL.Payload.for_deck(script, deck, s.spec,
                                             dpi=120, vector=True)))
    prs.save(str(deck))

    assert PL.refresh_deck(deck).updated
    again = Presentation(str(deck))
    shp = again.slides[0].shapes[0]
    assert V.has_svg(shp), "재생성 후 벡터가 래스터로 떨어졌습니다"
